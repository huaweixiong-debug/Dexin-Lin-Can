#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
读取PLC交互点位Excel文件
"""

import sys
import os

def read_excel_file():
    """读取Excel文件"""
    excel_file = "德鑫老化PLC交互点位.xlsx"
    
    if not os.path.exists(excel_file):
        print(f"文件不存在: {excel_file}")
        return
    
    print(f"读取文件: {excel_file}")
    print("=" * 80)
    
    # 尝试使用openpyxl
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        ws = wb.active
        
        print(f"工作表名称: {ws.title}")
        print(f"行数: {ws.max_row}, 列数: {ws.max_column}")
        print("-" * 80)
        
        # 读取前20行数据
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
            # 过滤空行
            if any(cell is not None for cell in row):
                print(f"行 {i+1}: {row}")
        
        print("-" * 80)
        
        # 查找关键信息
        print("查找PLC点位信息:")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)):
            if any(cell and isinstance(cell, str) and ('M' in str(cell) or 'Q' in str(cell) or 'I' in str(cell)) for cell in row):
                print(f"行 {i+1}: {row}")
        
    except ImportError:
        print("未安装openpyxl，尝试使用xlrd")
        try:
            import xlrd
            wb = xlrd.open_workbook(excel_file)
            ws = wb.sheet_by_index(0)
            
            print(f"工作表名称: {ws.name}")
            print(f"行数: {ws.nrows}, 列数: {ws.ncols}")
            print("-" * 80)
            
            # 读取前20行数据
            for i in range(min(20, ws.nrows)):
                row = ws.row_values(i)
                # 过滤空行
                if any(cell for cell in row):
                    print(f"行 {i+1}: {row}")
            
            print("-" * 80)
            
            # 查找关键信息
            print("查找PLC点位信息:")
            for i in range(ws.nrows):
                row = ws.row_values(i)
                if any(cell and isinstance(cell, str) and ('M' in str(cell) or 'Q' in str(cell) or 'I' in str(cell)) for cell in row):
                    print(f"行 {i+1}: {row}")
                    
        except ImportError:
            print("未安装openpyxl和xlrd，无法读取Excel文件")
            print("请安装: pip install openpyxl")
            return

if __name__ == "__main__":
    read_excel_file()